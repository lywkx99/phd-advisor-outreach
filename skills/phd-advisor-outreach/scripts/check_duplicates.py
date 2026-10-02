#!/usr/bin/env python3
"""Read-only advisor contact evidence search; never establishes sending status.

Names are aliases to search, not proof of identity. Institutions only annotate
name/email hits. Only explicit application-list email entries are promoted to
confirmed_application_list; that label does not establish an email was sent.
"""
import argparse
import html
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

EMAIL = re.compile(r"[a-z0-9.!#$%&'*+/=?^_`{|}~-]+@[a-z0-9](?:[a-z0-9.-]*[a-z0-9])?\.[a-z]{2,}")
LABELS = {
    "name": {"导师", "导师姓名", "姓名", "name", "pi"},
    "email": {"邮箱", "电子邮箱", "邮箱地址", "email", "e-mail"},
    "institution": {"单位", "学校", "机构", "院校", "institution", "university"},
}


def norm(value):
    value = unicodedata.normalize("NFKC", html.unescape(str(value)))
    return re.sub(r"\\([@_])", r"\1", value).casefold()


def emails(value):
    return sorted(set(EMAIL.findall(norm(value))))


def names_in(value, names):
    value = norm(value)
    return [n for n in names if (re.search(r"(?<![a-z0-9])" + re.escape(n) + r"(?![a-z0-9])", value)
                               if re.fullmatch(r"[a-z0-9 .'-]+", n) else n in value)]


class Search:
    def __init__(self, args):
        self.names = list(dict.fromkeys(norm(n).strip() for n in args.name if n.strip()))
        self.emails = list(dict.fromkeys(e for raw in args.email for e in emails(raw)))
        self.institution = norm(args.institution).strip() if args.institution else None
        self.result = {"query": {"names": self.names, "emails": self.emails,
                                 "institution": self.institution},
                       "scope": [], "matches": [], "warnings": [],
                       "interpretation": "仅检索本地记录；姓名命中需核对身份，草稿及表格不证明发送；无命中仅表示检索范围内未发现记录。"}
        for raw in args.email:
            if not emails(raw):
                self.warn("invalid_email", f"无法解析查询邮箱: {raw}")

    def warn(self, code, message, **location):
        warning = {"code": code, "message": message, **location}
        if warning not in self.result["warnings"]:
            self.result["warnings"].append(warning)

    def hits(self, value):
        return {"names": names_in(value, self.names),
                "emails": sorted(set(emails(value)) & set(self.emails))}

    def add(self, path, kind, value, reasons, **location):
        hit = self.hits(value)
        if not any(hit.values()):
            return
        context = re.sub(r"\s+", " ", str(value)).strip()
        self.result["matches"].append({"file": str(path), **location, "kind": kind,
                                       "matched": hit, "reasons": reasons,
                                       "context": context[:600] + ("…" if len(context) > 600 else "")})

    def check_identity(self, path, fields, **location):
        problems = []
        field_emails = emails(fields.get("email", ""))
        name_match = names_in(fields.get("name", ""), self.names)
        email_match = set(field_emails) & set(self.emails)
        if name_match and field_emails and self.emails and not email_match:
            problems.append("姓名命中但身份字段邮箱不同；不可判定为同一人")
            self.warn("name_email_conflict", problems[-1], file=str(path), **location)
        if email_match and self.names and fields.get("name") and not name_match:
            problems.append("邮箱命中但身份字段姓名与查询不符；需核对身份")
            self.warn("email_name_conflict", problems[-1], file=str(path), **location)
        institution = norm(fields.get("institution", "")).strip()
        if (name_match or email_match) and self.institution and institution:
            if self.institution not in institution and institution not in self.institution:
                problems.append("单位文本与查询不同（可能为别名）；单位仅作身份核对上下文")
                self.warn("institution_text_mismatch", problems[-1], file=str(path), **location)
        return problems

    def read_text(self, path):
        try:
            raw = path.read_bytes()
            try:
                return raw.decode("utf-8-sig").splitlines()
            except UnicodeDecodeError:
                self.warn("encoding_fallback", "非UTF-8文件，尝试GB18030", file=str(path))
                return raw.decode("gb18030").splitlines()
        except (OSError, UnicodeError) as exc:
            self.warn("read_failed", str(exc), file=str(path))
            return None

    def scope(self, kind, path):
        path = Path(path).expanduser().absolute()
        item = {"kind": kind, "path": str(path), "exists": path.exists(), "files_read": 0}
        self.result["scope"].append(item)
        if not item["exists"]:
            self.warn("missing_path", "检索路径不存在", file=str(path))
        return path, item

    @staticmethod
    def columns(cells):
        mapping = {}
        for i, cell in enumerate(cells):
            label = norm(cell).strip().strip("* ")
            for key, labels in LABELS.items():
                if label in labels:
                    mapping[key] = i
        return mapping if "name" in mapping or "email" in mapping else {}

    def memory(self, raw_path):
        path, scope = self.scope("memory", raw_path)
        if not scope["exists"]:
            return
        lines = self.read_text(path)
        scope["files_read"] = int(lines is not None)
        if lines is None:
            return
        application = False
        fenced = False
        columns = {}
        for i, line in enumerate(lines, 1):
            clean = norm(line).strip()
            heading = clean.startswith("#") or bool(re.match(r"^\*\*[^*]+\*\*", clean))
            if heading:
                application = bool(re.search(r"已(?:申请|发送|联系).*(?:邮箱|名单|清单|导师)|已(?:申请|发送|联系)\s*[:：*]*$", clean)
                                   and not re.search(r"未|查重|草稿|起草", clean))
                columns = {}
            if clean.startswith("```"):
                if fenced:
                    application = False
                fenced = not fenced
                continue
            cells = [c.strip() for c in line.strip().strip("|").split("|")] if clean.startswith("|") else []
            new_columns = self.columns(cells)
            if new_columns:
                columns = new_columns
                continue
            hit = self.hits(line)
            if not any(hit.values()):
                continue
            kind = "memory_context_match"
            reasons = ["姓名或邮箱出现在账本文本；按原文核对身份与状态"]
            fields = {key: cells[col] for key, col in columns.items() if col < len(cells)} if cells else {}
            identity = self.hits(" | ".join(fields.get(k, "") for k in ("name", "email")))
            problems = self.check_identity(path, fields, line=i) if fields else []
            if fields and any(identity.values()):
                kind = "memory_record_candidate"
                reasons = ["账本表格身份列命中；姓名相同不能自动认定为同一人"]
                if re.search(r"起草|草稿|邮件已写", clean):
                    kind = "memory_draft_record"
                    reasons.append("原文记载草稿；发送状态未由本工具核实")
            # Only a direct identity email or a bare email list entry can prove list membership.
            bare = re.sub(r"^[\s>*+\-\d.)]+", "", clean).strip("` ")
            direct_email = identity["emails"] if fields else (hit["emails"] if EMAIL.fullmatch(bare) else [])
            if application and direct_email and not problems:
                kind = "confirmed_application_list"
                reasons = ["查询邮箱精确命中明确的已申请名单；仅确认名单收录，发送时间及渠道未知"]
            if problems:
                kind = "identity_conflict_candidate"
                reasons += problems
            self.add(path, kind, line, reasons, line=i)

    def drafts(self, raw_path):
        root, scope = self.scope("drafts", raw_path)
        if not scope["exists"]:
            return
        try:
            files = [root] if root.is_file() else sorted(root.rglob("*"))
        except OSError as exc:
            self.warn("read_failed", str(exc), file=str(root))
            return
        supported = {".md", ".txt", ".eml", ".markdown"}
        scope["extensions"] = sorted(supported)
        for path in files:
            if not path.is_file() or path.suffix.casefold() not in supported:
                continue
            lines = self.read_text(path)
            scope["files_read"] += int(lines is not None)
            if lines is None:
                continue
            header_rows = set()
            recipient_emails = []
            recipient_name = ""
            title = ""
            for i, line in enumerate(lines[:40], 1):
                clean = norm(line).strip().replace("**", "")
                if not clean:
                    continue
                if re.match(r"^(?:邮件主题|主题|subject|尊敬的|dear\b|您好)", clean):
                    break
                if i == 1 and clean.startswith("#"):
                    title = line
                    header_rows.add(i)
                elif re.match(r"^(?:收件邮箱|收件人邮箱|收件人|to|recipient(?: email)?)\s*[:：]", clean):
                    header_rows.add(i)
                    recipient_emails.extend(emails(clean))
                    if re.match(r"^(?:收件人|to|recipient)\s*[:：]", clean):
                        recipient_name = EMAIL.sub("", clean.split(":", 1)[-1]).strip(" <>[](),;\"'")
                elif not re.match(r"^(?:查重|单位|学校|状态|from|date|cc|bcc|reply-to|message-id|mime-version|content-type)\s*[:：]", clean):
                    # An unrecognised prose line ends the header, even without a subject.
                    break
            fields = {"name": recipient_name or title, "email": " ".join(recipient_emails)}
            owner_names = names_in(recipient_name or title or path.stem, self.names)
            exact_recipient = set(recipient_emails) & set(self.emails)
            problems = self.check_identity(path, fields)
            own_kind = "draft_recipient_email_match" if exact_recipient else "draft_name_candidate"
            own_reasons = (["草稿头收件邮箱精确匹配；为本人草稿候选，发送状态未核实"] if exact_recipient else
                           ["仅草稿标题、收件人或文件名姓名匹配；同名不能确认身份或发送状态"])
            if problems:
                own_kind = "identity_conflict_candidate"
                own_reasons += problems
            filename_hit = self.hits(path.name)
            if any(filename_hit.values()):
                self.add(path, own_kind if owner_names or exact_recipient else "filename_candidate",
                         path.name, own_reasons, location="filename", recipient_emails=sorted(set(recipient_emails)))
            for i, line in enumerate(lines, 1):
                is_owner = i in header_rows and (owner_names or exact_recipient)
                self.add(path, own_kind if is_owner else "literature_or_context_match", line,
                         own_reasons if is_owner else ["正文、文献或其他上下文提及；不能据此认定为该导师的草稿或联系记录"],
                         line=i, **({"recipient_emails": sorted(set(recipient_emails))} if is_owner else {}))

    def table(self, raw_path):
        path, scope = self.scope("table", raw_path)
        if not scope["exists"]:
            return
        try:
            import openpyxl
        except ImportError:
            self.warn("optional_dependency_missing", "读取xlsx需安装openpyxl；该表未检索", file=str(path))
            return
        workbook = None
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            scope["files_read"] = 1
            scope["sheets"] = workbook.sheetnames
            for sheet in workbook:
                columns = {}
                for row, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [str(v) if v is not None else "" for v in values]
                    new_columns = self.columns(cells)
                    if new_columns:
                        columns = new_columns
                        continue
                    value = " | ".join(cells)
                    if not any(self.hits(value).values()):
                        continue
                    fields = {k: cells[c] for k, c in columns.items() if c < len(cells)}
                    identity = self.hits(" | ".join(fields.get(k, "") for k in ("name", "email")))
                    problems = self.check_identity(path, fields, sheet=sheet.title, row=row)
                    kind = "table_record_candidate" if any(identity.values()) else "table_context_match"
                    self.add(path, "identity_conflict_candidate" if problems else kind, value,
                             ["表格出现姓名或邮箱不等于已联系；需另核实身份及发送证据"] + problems,
                             sheet=sheet.title, row=row)
        except Exception as exc:
            self.warn("table_read_failed", str(exc), file=str(path))
        finally:
            if workbook is not None:
                workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--name", action="append", default=[], help="姓名/姓名别名，可重复")
    parser.add_argument("--email", action="append", default=[], help="邮箱，可重复；自动规范化")
    parser.add_argument("--institution", help="单位，仅用于命中后的上下文核对")
    parser.add_argument("--ledger", "--memory", dest="memory", action="append", default=[],
                        help="联系账本或交接文件（Markdown），可重复")
    parser.add_argument("--drafts", action="append", default=[], help="草稿文件或目录，可重复")
    parser.add_argument("--table", action="append", default=[], help="可选xlsx路径，可重复")
    args = parser.parse_args()
    if not args.name and not args.email:
        parser.error("至少提供一个 --name 或 --email；机构不是独立检索条件")
    if not (args.memory or args.drafts or args.table):
        parser.error("至少提供一个检索范围：--ledger、--drafts 或 --table")
    search = Search(args)
    if search.names or search.emails:
        for path in args.memory:
            search.memory(path)
        for path in args.drafts:
            search.drafts(path)
        for path in args.table:
            search.table(path)
    search.result["counts"] = dict(Counter(m["kind"] for m in search.result["matches"]))
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(search.result, ensure_ascii=False, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

